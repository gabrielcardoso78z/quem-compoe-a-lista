#!/usr/bin/env python3
"""2a rodada: acordo e kappa de Cohen (categorias completas e grande grupo), matriz de confusao por grande grupo,
subamostra de grafias novas (nao vistas na 1a rodada) e, se houver dois codificadores, kappa entre eles.
Confere que codificar_profissoes.py nao mudou desde o sorteio (REGRAS_CONGELADAS.txt).
Uso: python3 calcular_kappa_r2.py codificacao_r2_Gabriel.xlsx [codificacao_r2_professor.xlsx]"""
import csv, hashlib, os, sys
from collections import Counter
AQUI = os.path.dirname(os.path.abspath(__file__)); BASE = os.path.dirname(os.path.dirname(AQUI))
h = hashlib.sha256(open(os.path.join(BASE, "codificar_profissoes.py"), "rb").read()).hexdigest()
if h not in open(os.path.join(AQUI, "REGRAS_CONGELADAS.txt")).read(): sys.exit("ATENCAO: as regras mudaram depois do sorteio; a validacao nao vale.")
gab = {}; nova = {}
for r in csv.DictReader(open(os.path.join(AQUI, "gabarito_r2_NAO_ENVIAR.csv"), encoding="utf-8-sig"), delimiter=";"):
    gab[r["id"]] = r["categoria_script"]; nova[r["id"]] = r["grafia_ja_na_rodada1"] == "nao"
from openpyxl import load_workbook
def ler(arq):
    wb = load_workbook(arq, data_only=True)
    def cod(x):
        if x is None: return ""
        if isinstance(x, float) and x.is_integer(): x = int(x)
        return str(x).strip().upper()
    cats = {cod(r[0]): r[1] for r in wb["Categorias"].iter_rows(min_row=2, values_only=True) if r[0] is not None}
    return {str(int(r[0]) if isinstance(r[0], float) else r[0]): cats.get(cod(r[2]), "") for r in wb["Codificar"].iter_rows(min_row=2, values_only=True) if r[0] is not None}
def kappa(p):
    n = len(p); po = sum(a == b for a, b in p) / n; ca, cb = Counter(a for a, _ in p), Counter(b for _, b in p)
    pe = sum(ca[k] * cb[k] for k in ca) / n ** 2; return po, (po - pe) / (1 - pe), n
GG = lambda c: c[0] if c[:1].isdigit() else "fora"
cods = [ler(a) for a in sys.argv[1:]]
for arq, col in zip(sys.argv[1:], cods):
    ids = [i for i in gab if col.get(i)]
    pares = [(gab[i], col[i]) for i in ids]
    if not pares: print(f"{arq}: ainda vazia"); continue
    print(f"\n== {os.path.basename(arq)} x programa: {len(pares)} entradas codificadas ==")
    for nome, f in [("categorias completas", lambda c: c), ("grande grupo COD", GG)]:
        po, k, n = kappa([(f(a), f(b)) for a, b in pares]); print(f"{nome}: acordo {100*po:.1f}% ({sum(f(a)==f(b) for a,b in pares)}/{n}), kappa {k:.2f}")
    pn = [(gab[i], col[i]) for i in ids if nova[i]]
    if pn:
        po, k, n = kappa(pn); print(f"so grafias novas (nao vistas na 1a rodada): acordo {100*po:.1f}% ({n} entradas), kappa {k:.2f}")
    lab = sorted({GG(a) for a, _ in pares} | {GG(b) for _, b in pares})
    m = Counter((GG(a), GG(b)) for a, b in pares)
    print("matriz de confusao (linhas: programa; colunas: codificador):")
    print("      " + "".join(f"{l:>6}" for l in lab))
    for a in lab: print(f"{a:>6}" + "".join(f"{m[(a,b)]:>6}" for b in lab))
    print("divergencias:")
    for (a, b), n in Counter((a, b) for a, b in pares if a != b).most_common(): print(f"  {n:3}  {a}  ->  {b}")
if len(cods) == 2:
    p = [(cods[0][i], cods[1][i]) for i in gab if cods[0].get(i) and cods[1].get(i)]
    if p:
        po, k, n = kappa(p); print(f"\n== entre os dois codificadores: acordo {100*po:.1f}% ({n}), kappa {k:.2f} ==")
