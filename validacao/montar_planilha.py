#!/usr/bin/env python3
"""Monta codificacao_manual.xlsx (abre no Excel/Numbers) a partir de amostra_para_codificar.csv."""
import csv
from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Font, Alignment, PatternFill
CODIGOS = ["0","1","2","3","4","5","6","7","8","9","E","A","L","D","S","F","O"]
CATS = [
 ("0 Forcas armadas, policiais e bombeiros militares", "Militares das Forcas Armadas, policiais militares e bombeiros militares da ativa."),
 ("1 Diretores e gerentes", "Quem dirige, gerencia ou e dono de empresa ou orgao: empresario, gerente, diretor, coordenador, supervisor."),
 ("2 Profissionais das ciencias e intelectuais", "Ocupacoes que em regra exigem curso superior: advogado, medico, enfermeiro, professor, engenheiro, contador, psicologo, analista."),
 ("3 Tecnicos e profissionais de nivel medio", "Tecnicos e ocupacoes de nivel medio especializado: tecnico de enfermagem, tecnico em informatica, corretor, representante comercial."),
 ("4 Trabalhadores de apoio administrativo", "Escritorio e atendimento: auxiliar ou assistente administrativo, secretaria, recepcionista, bancario (escriturario), caixa."),
 ("5 Trabalhadores dos servicos e vendedores", "Servicos pessoais, protecao e comercio: vendedor, comerciante, cabeleireiro, cozinheiro, vigilante, porteiro."),
 ("6 Agropecuaria", "Agricultor, produtor rural, trabalhador rural, jardineiro."),
 ("7 Operarios e artesaos", "Oficios manuais qualificados: pedreiro, eletricista, mecanico, costureira, soldador, padeiro."),
 ("8 Operadores de maquinas e montadores", "Operadores de maquinas e condutores de veiculos: motorista, operador de maquinas, montador."),
 ("9 Ocupacoes elementares", "Tarefas simples: diarista, faxineiro, servicos gerais, ajudante, servente, empregada domestica."),
 ("Estudante ou estagiario", "Estudante, estagiario, bolsista, aprendiz (mesmo que diga o curso)."),
 ("Aposentado", "Aposentado, pensionista, inativo, da reserva (mesmo que diga a profissao anterior)."),
 ("Do lar", "Do lar, dona de casa."),
 ("Desempregado", "Desempregado, sem ocupacao."),
 ("Servidor publico sem ocupacao informada", "So diz que e servidor ou funcionario publico, sem o cargo."),
 ("Empregado sem ocupacao informada", "So diz funcionario, empregado, colaborador, CLT, sem a ocupacao."),
 ("Outros ou nao classificado", "Nao da para saber a ocupacao (ex.: autonomo, prestador de servicos) ou nao se encaixa em nada."),
]
wb = Workbook()
ins = wb.active; ins.title = "Instrucoes"
txt = [
 "COMO CODIFICAR (cerca de 1 hora)",
 "1. Va para a aba 'Codificar'. Para cada linha, leia a profissao e DIGITE na coluna C o CODIGO da categoria (0 a 9 para os grupos de ocupacao; E estudante; A aposentado; L do lar; D desempregado; S servidor sem cargo; F empregado sem ocupacao; O outros). A lista de codigos esta na aba 'Categorias'.",
 "2. Use a aba 'Categorias' como guia. Em caso de duvida sobre o grande grupo, consulte a estrutura oficial da COD do IBGE (Classificacao de Ocupacoes para Pesquisas Domiciliares).",
 "3. Codifique o que esta escrito, nao o que voce imagina. 'Servidor publico' sem cargo = 'Servidor publico sem ocupacao informada', mesmo que voce desconfie do cargo.",
 "4. Duas ocupacoes na mesma linha: use a primeira.",
 "5. Aposentado ou estudante prevalece sobre a profissao (ex.: 'professor aposentado' = Aposentado; 'estudante de enfermagem' = Estudante).",
 "6. Se nao der para decidir, escolha 'Outros ou nao classificado' e, se quiser, explique na coluna D.",
 "7. IMPORTANTE: nao abra o arquivo gabarito_NAO_ENVIAR.csv nem a planilha codificacao_profissoes.csv antes de terminar. A validacao so vale se voce codificar sem ver a classificacao automatica.",
 "8. Ao terminar, salve este arquivo com o mesmo nome e avise. O script calcular_kappa.py compara as suas escolhas com as do programa.",
]
for i, t in enumerate(txt, 1):
    c = ins.cell(row=i, column=1, value=t); c.alignment = Alignment(wrap_text=True)
    if i == 1: c.font = Font(bold=True, size=13)
ins.column_dimensions["A"].width = 120
cat = wb.create_sheet("Categorias")
cat.append(["Codigo", "Categoria", "O que entra (exemplos)"])
for cod, (c, d) in zip(CODIGOS, CATS): cat.append([cod, c, d])
cat.column_dimensions["A"].width = 8; cat.column_dimensions["B"].width = 48; cat.column_dimensions["C"].width = 110
for c in cat[1]: c.font = Font(bold=True)
ws = wb.create_sheet("Codificar", 1)
ws.append(["id", "Profissao (como publicada)", "Codigo (veja a aba Categorias)", "Observacao (opcional)"])
for c in ws[1]: c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="DDDDDD")
rows = list(csv.DictReader(open("amostra_para_codificar.csv", encoding="utf-8-sig"), delimiter=";"))
for r in rows: ws.append([int(r["id"]), r["profissao (como publicada)"], "", ""])
dv = DataValidation(type="list", formula1='"' + ",".join(CODIGOS) + '"', allow_blank=True)
dv.error = "Use um dos codigos da aba Categorias"; dv.showErrorMessage = True
ws.add_data_validation(dv); dv.add(f"C2:C{len(rows)+1}")
ws.column_dimensions["A"].width = 6; ws.column_dimensions["B"].width = 50; ws.column_dimensions["C"].width = 30; ws.column_dimensions["D"].width = 40
ws.freeze_panes = "A2"
wb.active = 0
wb.save("codificacao_manual.xlsx"); print("ok", len(rows))
