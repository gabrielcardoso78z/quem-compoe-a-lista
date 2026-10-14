#!/usr/bin/env python3
"""Concordancia entre o script e o segundo codificador: % de acordo e kappa de Cohen
(categorias completas e so o grande grupo COD). Uso: python3 calcular_kappa.py amostra_preenchida.csv"""
import csv, sys, os
from collections import Counter
AQUI = os.path.dirname(os.path.abspath(__file__))
gab = {r["id"]: r["categoria_script"] for r in csv.DictReader(open(os.path.join(AQUI, "gabarito_NAO_ENVIAR.csv"), encoding="utf-8-sig"), delimiter=";")}
arq = sys.argv[1] if len(sys.argv) > 1 else os.path.join(AQUI, "codificacao_manual.xlsx")
if arq.endswith(".xlsx"):
    from openpyxl import load_workbook
    ws = load_workbook(arq, data_only=True)["Codificar"]
    cats = {str(r[0]).strip().upper(): r[1] for r in load_workbook(arq, data_only=True)["Categorias"].iter_rows(min_row=2, values_only=True) if r[0] is not None}
    def cod(x):
        if x is None: return ""
        if isinstance(x, float) and x.is_integer(): x = int(x)   # Google Planilhas grava 2 como 2.0
        return str(x).strip().upper()
    cats = {cod(c): n for c, n in cats.items()}
    col = {str(int(r[0]) if isinstance(r[0], float) else r[0]): cats.get(cod(r[2]), cod(r[2])) for r in ws.iter_rows(min_row=2, values_only=True) if r[0] is not None}
else:
    col = {r["id"]: r["categoria (escolha uma da aba de categorias)"].strip() for r in csv.DictReader(open(arq, encoding="utf-8-sig"), delimiter=";")}
def kappa(pares):
    n = len(pares); po = sum(a == b for a, b in pares) / n
    ca, cb = Counter(a for a, _ in pares), Counter(b for _, b in pares)
    pe = sum(ca[k] * cb[k] for k in ca) / n ** 2
    return po, (po - pe) / (1 - pe)
pares = [(gab[i], col[i]) for i in gab if col.get(i)]
if not pares: sys.exit("A coluna de categoria ainda esta vazia.")
for nome, f in [("categorias completas", lambda c: c), ("so grande grupo COD (demais = fora)", lambda c: c[0] if c[:1].isdigit() else "fora")]:
    po, k = kappa([(f(a), f(b)) for a, b in pares])
    print(f"{nome}: n={len(pares)}, acordo {100*po:.1f}%, kappa {k:.2f}")
print("Divergencias (script -> colega):")
for (a, b), n in Counter((a, b) for a, b in pares if a != b).most_common(): print(f"  {n:3}  {a}  ->  {b}")
