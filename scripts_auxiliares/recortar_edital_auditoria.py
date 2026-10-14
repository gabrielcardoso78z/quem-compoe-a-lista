"""Recorta do PDF do Diario as paginas de um edital da amostra da auditoria e registra resultados.
Uso:  python3 recortar_edital_auditoria.py N            -> recorta a linha N da planilha (1 = primeira)
      python3 recortar_edital_auditoria.py --ok N       -> marca a linha N como conferida, tudo ok
O recorte vai para auditoria/pdfs/ (fica so no computador do pesquisador). Nao imprime nomes."""
import os, re, subprocess, sys, unicodedata
from openpyxl import load_workbook
BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(BASE, "auditoria", "auditoria_editais.xlsx")

def slug(s):
    s = unicodedata.normalize("NFD", s); s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return re.sub(r"[^A-Za-z0-9]+", "_", s).strip("_")

def nums_pagina(pdf, p):
    t = subprocess.run(["pdftotext", "-f", str(p), "-l", str(p), "-raw", pdf, "-"], capture_output=True, text=True).stdout
    return [int(x) for x in re.findall(r"(?:^|\s)(\d{1,4})\s*[-.)]", t)]

wb = load_workbook(XLSX); ws = wb["Auditoria"]
if sys.argv[1] == "--ok":
    n = int(sys.argv[2]); r = n + 1
    conv = ws.cell(r, 1).value == "C"
    for col in range(21, 26):
        ws.cell(r, col).value = "ok" if (not conv or col == 21 or col == 24) else "não se aplica"
    wb.save(XLSX); print("linha", n, "marcada ok"); sys.exit()

n = int(sys.argv[1]); r = n + 1
v = [ws.cell(r, c).value for c in range(1, 21)]
grupo, com, data, ed, arq, pag = v[0], v[3], v[4], v[5], v[6], str(v[7])
pdf = os.path.join(BASE, "dje", arq)
ini = int(re.match(r"\d+", pag).group()); total = int(subprocess.run(["pdfinfo", pdf], capture_output=True, text=True).stdout.split("Pages:")[1].split()[0])
fim = ini
if grupo != "C":
    ult = int(v[12] or 0)
    p = ini
    while p <= min(total, ini + 40):
        ns = nums_pagina(pdf, p)
        fim = p
        if ult and ult in ns and any(ult - 15 <= x < ult for x in ns): break
        p += 1
    if fim - ini > 30: fim = ini + 2
    ini = max(1, ini - (1 if "anterior" in pag else 0))
else:
    fim = min(total, ini + 1)
os.makedirs(os.path.join(BASE, "auditoria", "pdfs"), exist_ok=True)
saida = os.path.join(BASE, "auditoria", "pdfs", f"{grupo}{n:02d}_{slug(com)}_{data}_ed{ed}_p{ini}-{fim}.pdf")
subprocess.run(["qpdf", pdf, "--pages", ".", f"{ini}-{fim}", "--", saida], check=True)
ws.cell(r, 8).value = f"{ini}-{fim}" if ini != fim else str(ini)
wb.save(XLSX)
falt = []
if grupo != "C":
    import csv
    sq = {int(e["seq"]) for e in csv.DictReader(open(os.path.join(BASE, "data", "comarcas_pr_entradas.csv"), encoding="utf-8"))
          if e["comarca"] == com and e["data_dje"] == data and e["edicao"] == str(ed)}
    if sq: falt = [x for x in range(min(sq), max(sq) + 1) if x not in sq]
print(saida); print("números não lidos:", falt[:40], "(total", len(falt), ")"); print("linha", n, "|", grupo, "|", com, data, "| páginas", ini, "a", fim, "| caráter", v[8], v[9], "| lidas", v[10], "| 1º", v[11], "último", v[12], "| ausentes", v[13], "| posições", v[14], v[16], v[18])
